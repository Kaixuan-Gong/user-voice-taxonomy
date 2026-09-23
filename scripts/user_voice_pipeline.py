#!/usr/bin/env python3
"""Deterministic pipeline for user-voice taxonomy discovery and hierarchical labeling.

The executing AI supplies semantic_features.jsonl and taxonomy.json according to
references/contracts.md. This script owns ingestion, embedding,
clustering, taxonomy validation, Laya inference, review gates, metrics and xlsx.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import os
import re
import sys
import time
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

import numpy as np
import pandas as pd

NOISE = {"", "无", "暂无", "不知道", "没问题", "ok", "好的", "收到", "谢谢"}
DEFAULT_THRESHOLDS = [0.62, 0.66, 0.70, 0.72, 0.74, 0.76, 0.78, 0.80]
REQUIRED_FEATURE_FIELDS = {
    "record_id", "key_points", "problem_object", "user_request", "cause",
    "emotion", "severity", "evidence_spans", "multi_intent", "issue_spans",
}
REQUIRED_LABEL_FIELDS = {
    "label_id", "parent_id", "level", "name", "definition", "include",
    "exclude", "positive_examples", "negative_examples", "business_action", "status",
}


def dump_json(path: Path, obj: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=2), encoding="utf-8")


def load_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def read_jsonl(path: Path) -> list[dict[str, Any]]:
    out = []
    with path.open("r", encoding="utf-8") as f:
        for line_no, line in enumerate(f, 1):
            line = line.strip()
            if not line:
                continue
            try:
                out.append(json.loads(line))
            except json.JSONDecodeError as e:
                raise ValueError(f"{path}:{line_no} JSON 格式错误: {e}") from e
    return out


def write_jsonl(path: Path, rows: Iterable[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as f:
        for row in rows:
            f.write(json.dumps(row, ensure_ascii=False) + "\n")


def normalize_text(value: Any) -> str:
    text = "" if value is None or (isinstance(value, float) and math.isnan(value)) else str(value)
    text = text.replace("\u3000", " ")
    text = re.sub(r"\s+", " ", text).strip()
    text = re.sub(r"([!！?？。,.，])\1+", r"\1", text)
    return text


def read_input(path: Path, text_col: str, id_col: str | None, sheet_name: str | None) -> pd.DataFrame:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        df = pd.read_csv(path, dtype=str, keep_default_na=False)
    elif suffix in {".xlsx", ".xls"}:
        df = pd.read_excel(path, sheet_name=sheet_name or 0, dtype=str, keep_default_na=False)
    elif suffix == ".parquet":
        df = pd.read_parquet(path)
    elif suffix in {".jsonl", ".ndjson"}:
        df = pd.DataFrame(read_jsonl(path))
    elif suffix == ".json":
        obj = load_json(path)
        df = pd.DataFrame(obj if isinstance(obj, list) else obj.get("data", []))
    else:
        raise ValueError(f"不支持的输入格式: {suffix}")
    if text_col not in df.columns:
        raise ValueError(f"找不到原声列 {text_col!r}，现有列: {list(df.columns)}")
    if id_col and id_col not in df.columns:
        raise ValueError(f"找不到 ID 列 {id_col!r}")
    return df


def prepare_records(args: argparse.Namespace) -> dict[str, Any]:
    src = Path(args.input).expanduser().resolve()
    outdir = Path(args.output_dir).expanduser().resolve()
    outdir.mkdir(parents=True, exist_ok=True)
    df = read_input(src, args.text_col, args.id_col, args.sheet_name)
    raw_count = len(df)
    records, rejected, seen = [], [], {}
    for idx, row in df.iterrows():
        rid = str(row[args.id_col]) if args.id_col else str(idx + 1)
        raw = normalize_text(row[args.text_col])
        if not raw or raw.lower() in NOISE or len(raw) < args.min_chars:
            rejected.append({"record_id": rid, "raw_text": raw, "reason": "empty_or_noise"})
            continue
        fp = hashlib.sha256(raw.encode("utf-8")).hexdigest()[:16]
        if fp in seen:
            rejected.append({"record_id": rid, "raw_text": raw, "reason": f"duplicate_of:{seen[fp]}"})
            continue
        seen[fp] = rid
        context = {str(c): normalize_text(row[c]) for c in (args.context_cols or []) if c in df.columns}
        records.append({
            "record_id": rid,
            "source_row": int(idx) + 2,
            "raw_text": raw,
            "cleaned_text": raw,
            "context": context,
            "fingerprint": fp,
        })
    write_jsonl(outdir / "normalized_records.jsonl", records)
    write_jsonl(outdir / "rejected_records.jsonl", rejected)
    batches = []
    for start in range(0, len(records), args.batch_size):
        chunk = records[start:start + args.batch_size]
        batches.append({"batch_id": len(batches) + 1, "records": [
            {"record_id": r["record_id"], "text": r["cleaned_text"], "context": r["context"]} for r in chunk
        ]})
    write_jsonl(outdir / "ai_feature_batches.jsonl", batches)
    profile = {
        "input": str(src), "text_col": args.text_col, "id_col": args.id_col,
        "raw_count": raw_count, "valid_count": len(records), "rejected_count": len(rejected),
        "duplicate_count": sum(str(x["reason"]).startswith("duplicate_of") for x in rejected),
        "batch_count": len(batches), "generated_at": time.strftime("%Y-%m-%d %H:%M:%S"),
    }
    dump_json(outdir / "data_profile.json", profile)
    return profile


def validate_features(records: list[dict[str, Any]], features: list[dict[str, Any]]) -> list[str]:
    errors, by_id = [], {str(r["record_id"]): r for r in records}
    feature_ids = set()
    for i, f in enumerate(features, 1):
        missing = REQUIRED_FEATURE_FIELDS - set(f)
        if missing:
            errors.append(f"feature row {i} 缺字段: {sorted(missing)}")
            continue
        rid = str(f["record_id"])
        feature_ids.add(rid)
        if rid not in by_id:
            errors.append(f"feature row {i} record_id 不存在: {rid}")
            continue
        source = by_id[rid]["cleaned_text"]
        for span in f.get("evidence_spans", []):
            if span and span not in source:
                errors.append(f"record {rid} 证据片段不在原文: {span}")
        for issue in f.get("issue_spans", []):
            span = issue.get("evidence", "") if isinstance(issue, dict) else str(issue)
            if span and span not in source:
                errors.append(f"record {rid} issue 证据不在原文: {span}")
    missing_ids = set(by_id) - feature_ids
    if missing_ids:
        errors.append(f"缺少 {len(missing_ids)} 条特征记录: {sorted(missing_ids)[:10]}")
    return errors


def load_embedder(model_path: str):
    from sentence_transformers import SentenceTransformer
    device = "mps" if sys.platform == "darwin" else None
    try:
        return SentenceTransformer(model_path, device=device)
    except Exception:
        return SentenceTransformer(model_path, device="cpu")


def feature_text(record: dict[str, Any], feature: dict[str, Any]) -> str:
    points = feature.get("key_points") or []
    issues = feature.get("issue_spans") or []
    issue_text = "；".join(
        (x.get("issue") or x.get("evidence") or "") if isinstance(x, dict) else str(x) for x in issues
    )
    return "；".join(x for x in [record["cleaned_text"], "；".join(map(str, points)), issue_text] if x)


def cluster_for_threshold(emb: np.ndarray, threshold: float, min_cluster_size: int) -> np.ndarray:
    from sklearn.cluster import AgglomerativeClustering
    if len(emb) == 1:
        return np.array([0])
    labels = AgglomerativeClustering(
        n_clusters=None, metric="cosine", linkage="average", distance_threshold=1.0 - threshold,
    ).fit_predict(emb)
    counts = Counter(labels)
    remap, nxt = {}, 0
    out = []
    for label in labels:
        if counts[label] < min_cluster_size:
            out.append(-1)
        else:
            if label not in remap:
                remap[label] = nxt; nxt += 1
            out.append(remap[label])
    return np.asarray(out)


def silhouette_safe(emb: np.ndarray, labels: np.ndarray) -> float | None:
    from sklearn.metrics import silhouette_score
    mask = labels >= 0
    valid = labels[mask]
    if mask.sum() < 3 or len(set(valid)) < 2 or len(set(valid)) >= mask.sum():
        return None
    return float(silhouette_score(emb[mask], valid, metric="cosine"))


def representative_indices(emb: np.ndarray, labels: np.ndarray, n: int = 5) -> dict[int, list[int]]:
    out = {}
    for label in sorted(set(labels) - {-1}):
        ids = np.where(labels == label)[0]
        centroid = emb[ids].mean(axis=0)
        centroid /= np.linalg.norm(centroid) + 1e-12
        sims = emb[ids] @ centroid
        out[int(label)] = ids[np.argsort(-sims)[:n]].tolist()
    return out


def choose_threshold(diagnostics: list[dict[str, Any]], default: float = 0.72) -> float:
    """Pick a defensible threshold while penalizing fragmentation and giant/noise clusters."""
    candidates = []
    for d in diagnostics:
        sil = d["silhouette_cosine"]
        if sil is None or d["cluster_count"] < 2:
            continue
        noise_penalty = max(0.0, d["noise_rate"] - 0.35) * 1.2
        giant_penalty = max(0.0, d["largest_cluster_share"] - 0.45) * 1.2
        fragmentation_penalty = max(0.0, d["cluster_count"] - 50) * 0.002
        score = float(sil) - noise_penalty - giant_penalty - fragmentation_penalty
        candidates.append((score, -abs(d["similarity_threshold"] - default), d["similarity_threshold"]))
    return float(max(candidates)[2]) if candidates else float(default)


def cluster_records(args: argparse.Namespace) -> dict[str, Any]:
    work = Path(args.work_dir).expanduser().resolve()
    records = read_jsonl(work / "normalized_records.jsonl")
    features = read_jsonl(Path(args.features).expanduser().resolve())
    errors = validate_features(records, features)
    if errors:
        dump_json(work / "feature_validation_errors.json", errors)
        raise ValueError(f"语义特征校验失败，共 {len(errors)} 项，详见 feature_validation_errors.json")
    by_feature = {str(x["record_id"]): x for x in features}
    texts = [feature_text(r, by_feature[str(r["record_id"])]) for r in records]
    model = load_embedder(args.embedding_model)
    emb = model.encode(texts, batch_size=args.embedding_batch_size, normalize_embeddings=True, show_progress_bar=True)
    np.save(work / "embeddings.npy", emb)
    diagnostics = []
    for threshold in args.threshold_grid:
        labels = cluster_for_threshold(emb, threshold, args.min_cluster_size)
        diagnostics.append({
            "similarity_threshold": threshold,
            "cluster_count": len(set(labels) - {-1}),
            "noise_count": int((labels < 0).sum()),
            "noise_rate": round(float((labels < 0).mean()), 4),
            "largest_cluster_share": round(max(Counter(labels[labels >= 0]).values(), default=0) / max(len(labels), 1), 4),
            "silhouette_cosine": None if silhouette_safe(emb, labels) is None else round(silhouette_safe(emb, labels), 4),
        })
    chosen_threshold = args.similarity_threshold if args.similarity_threshold is not None else choose_threshold(diagnostics)
    for d in diagnostics:
        d["selected"] = abs(float(d["similarity_threshold"]) - float(chosen_threshold)) < 1e-9
    labels = cluster_for_threshold(emb, chosen_threshold, args.min_cluster_size)
    reps = representative_indices(emb, labels)
    cluster_rows, summaries = [], []
    for i, (record, feature, label) in enumerate(zip(records, features, labels)):
        cluster_rows.append({
            **record, **{k: feature.get(k) for k in REQUIRED_FEATURE_FIELDS if k != "record_id"},
            "cluster_id": "noise" if label < 0 else f"C{int(label):04d}",
        })
    for label in sorted(set(labels)):
        ids = np.where(labels == label)[0]
        summaries.append({
            "cluster_id": "noise" if label < 0 else f"C{int(label):04d}",
            "size": int(len(ids)),
            "share": round(len(ids) / max(len(labels), 1), 4),
            "representative_samples": [records[i]["cleaned_text"] for i in (reps.get(int(label), ids[:5].tolist()))],
            "key_points": Counter(p for i in ids for p in (features[i].get("key_points") or [])).most_common(12),
            "problem_objects": Counter(str(features[i].get("problem_object", "")) for i in ids).most_common(8),
        })
    write_jsonl(work / "clustered_records.jsonl", cluster_rows)
    dump_json(work / "cluster_summaries.json", summaries)
    dump_json(work / "threshold_diagnostics.json", diagnostics)
    dump_json(work / "cluster_manifest.json", {
        "embedding_model": args.embedding_model,
        "selected_similarity_threshold": chosen_threshold,
        "threshold_mode": "manual" if args.similarity_threshold is not None else "auto",
        "min_cluster_size": args.min_cluster_size,
    })
    return {"record_count": len(records), "cluster_count": len(set(labels) - {-1}), "noise_count": int((labels < 0).sum()), "selected_similarity_threshold": chosen_threshold}


def normalize_name(name: str) -> str:
    return re.sub(r"[\s_\-—/]+", "", str(name)).lower()


def validate_taxonomy_obj(taxonomy: dict[str, Any], embedding_model: str | None = None,
                          sibling_similarity_threshold: float = 0.86,
                          max_children: int = 20) -> dict[str, Any]:
    labels = taxonomy.get("labels", [])
    errors, warnings = [], []
    ids, by_parent = {}, defaultdict(list)
    for idx, label in enumerate(labels, 1):
        miss = REQUIRED_LABEL_FIELDS - set(label)
        if miss:
            errors.append(f"label row {idx} 缺字段: {sorted(miss)}")
            continue
        lid = str(label["label_id"])
        if lid in ids:
            errors.append(f"label_id 重复: {lid}")
        ids[lid] = label
        by_parent[label.get("parent_id")].append(label)
        if int(label.get("level", 0)) not in {1, 2, 3}:
            errors.append(f"{lid} level 必须为 1/2/3")
        if not str(label.get("definition", "")).strip():
            errors.append(f"{lid} definition 为空")
        if not label.get("positive_examples"):
            warnings.append(f"{lid} 缺少正例")
        if not label.get("negative_examples"):
            warnings.append(f"{lid} 缺少反例")
        if not str(label.get("business_action", "")).strip():
            warnings.append(f"{lid} 缺少业务动作")
    for lid, label in ids.items():
        parent_id = label.get("parent_id")
        level = int(label.get("level", 0))
        if level == 1 and parent_id not in {None, "", "null"}:
            errors.append(f"{lid} 是一级标签但 parent_id 非空")
        if level > 1:
            if parent_id not in ids:
                errors.append(f"{lid} 父标签不存在: {parent_id}")
            elif int(ids[parent_id]["level"]) != level - 1:
                errors.append(f"{lid} 与父标签层级不连续")
    for parent, children in by_parent.items():
        if len(children) > max_children:
            warnings.append(f"父节点 {parent} 有 {len(children)} 个子标签，超过建议上限 {max_children}")
        names = defaultdict(list)
        for x in children:
            names[normalize_name(x["name"])].append(x["label_id"])
        for name, dup_ids in names.items():
            if len(dup_ids) > 1:
                errors.append(f"同级标签名称重复 {name}: {dup_ids}")
    overlap_pairs = []
    if embedding_model and labels:
        model = load_embedder(embedding_model)
        for parent, children in by_parent.items():
            if len(children) < 2:
                continue
            texts = [f"{x['name']}：{x['definition']}；包含：{'；'.join(x.get('include', []))}" for x in children]
            emb = model.encode(texts, normalize_embeddings=True)
            sims = emb @ emb.T
            for i in range(len(children)):
                for j in range(i + 1, len(children)):
                    if sims[i, j] >= sibling_similarity_threshold:
                        pair = {"parent_id": parent, "label_a": children[i]["label_id"], "label_b": children[j]["label_id"], "similarity": round(float(sims[i, j]), 4)}
                        overlap_pairs.append(pair)
                        warnings.append(f"同级疑似重叠: {pair}")
    return {"valid": not errors, "errors": errors, "warnings": warnings, "overlap_pairs": overlap_pairs, "label_count": len(labels)}


def validate_taxonomy_cmd(args: argparse.Namespace) -> dict[str, Any]:
    taxonomy_path = Path(args.taxonomy).expanduser().resolve()
    taxonomy = load_json(taxonomy_path)
    result = validate_taxonomy_obj(taxonomy, args.embedding_model, args.sibling_similarity_threshold, args.max_children)
    output = Path(args.output or taxonomy_path.with_name("taxonomy_validation.json"))
    dump_json(output, result)
    if not result["valid"]:
        raise ValueError(f"标签体系校验失败，共 {len(result['errors'])} 个错误")
    return result


@dataclass
class Decision:
    label_id: str | None
    probability: float
    margin: float
    probabilities: dict[str, float]
    reason: str | None = None


class LayaClassifier:
    def __init__(self, model_path: str, subfolder: str = "multilingual"):
        import laya
        self.agent = laya.load(model_path, subfolder=subfolder)

    def choose(self, text: str, children: list[dict[str, Any]], parent_name: str) -> Decision:
        if len(children) == 1:
            return Decision(children[0]["label_id"], 1.0, 1.0, {children[0]["label_id"]: 1.0})
        criteria = {
            x["label_id"]: f"{x['name']}。定义：{x['definition']}。包含：{'；'.join(x.get('include', []))}。不包含：{'；'.join(x.get('exclude', []))}"
            for x in children
        }
        q = {"route": {"type": "choice", "instructions": f"判断这条用户原声在“{parent_name}”下最符合哪个标签。只根据标签定义选择。", "criteria": criteria}}
        ans = self.agent.predict(text, q)["answers"]["route"]
        probs = {str(k): float(v) for k, v in ans["probabilities"].items()}
        ordered = sorted(probs.values(), reverse=True)
        margin = ordered[0] - ordered[1] if len(ordered) > 1 else ordered[0]
        return Decision(str(ans["choice"]), probs[str(ans["choice"])], margin, probs)


class EmbeddingClassifier:
    def __init__(self, model_path: str):
        self.model = load_embedder(model_path)

    def choose(self, text: str, children: list[dict[str, Any]], parent_name: str) -> Decision:
        if len(children) == 1:
            return Decision(children[0]["label_id"], 1.0, 1.0, {children[0]["label_id"]: 1.0})
        defs = [f"{x['name']}：{x['definition']}。{'；'.join(x.get('include', []))}" for x in children]
        emb = self.model.encode([text] + defs, normalize_embeddings=True)
        sims = emb[1:] @ emb[0]
        exps = np.exp((sims - sims.max()) / 0.08); probs_arr = exps / exps.sum()
        probs = {x["label_id"]: float(p) for x, p in zip(children, probs_arr)}
        order = np.argsort(-probs_arr)
        best = int(order[0]); margin = float(probs_arr[order[0]] - probs_arr[order[1]])
        return Decision(children[best]["label_id"], float(probs_arr[best]), margin, probs)


def build_taxonomy(taxonomy: dict[str, Any]):
    labels = {x["label_id"]: x for x in taxonomy["labels"] if x.get("status", "active") == "active"}
    children = defaultdict(list)
    for label in labels.values():
        parent = label.get("parent_id") or None
        children[parent].append(label)
    return labels, children


def hierarchical_label(text: str, classifier: Any, labels: dict[str, dict[str, Any]], children: dict[Any, list[dict[str, Any]]], confidence_threshold: float, margin_threshold: float) -> dict[str, Any]:
    path, level_probs, full_probs, parent, reason = [], [], [], None, None
    while children.get(parent):
        opts = children[parent]
        parent_name = "全部问题" if parent is None else labels[parent]["name"]
        decision = classifier.choose(text, opts, parent_name)
        full_probs.append(decision.probabilities)
        if decision.label_id is None or decision.probability < confidence_threshold:
            reason = f"low_confidence:{decision.probability:.4f}"; break
        if decision.margin < margin_threshold:
            reason = f"small_margin:{decision.margin:.4f}"; break
        parent = decision.label_id
        path.append(parent)
        level_probs.append(decision.probability)
        if len(path) >= 3:
            break
    return {
        "label_ids": path,
        "label_path": [labels[x]["name"] for x in path],
        "probabilities": [round(x, 4) for x in level_probs],
        "confidence": round(min(level_probs), 4) if level_probs else 0.0,
        "all_probabilities": full_probs,
        "status": "review" if reason else "labeled",
        "review_reason": reason,
    }


def issue_texts(record: dict[str, Any], feature: dict[str, Any]) -> list[tuple[int, str, str]]:
    issues = feature.get("issue_spans") or []
    if not issues:
        return [(1, record["cleaned_text"], "")]
    out = []
    for i, issue in enumerate(issues, 1):
        if isinstance(issue, dict):
            evidence = str(issue.get("evidence", ""))
            issue_desc = str(issue.get("issue", ""))
            text = f"{evidence}。问题点：{issue_desc}" if evidence else record["cleaned_text"]
        else:
            evidence, text = str(issue), str(issue)
        out.append((i, text, evidence))
    return out


def make_distribution(tagged: list[dict[str, Any]]) -> list[dict[str, Any]]:
    total = len(tagged)
    counter = Counter(tuple(x.get("label_path", [])) if x.get("label_path") else ("待复核/未知",) for x in tagged)
    rows = []
    for path, count in counter.most_common():
        rows.append({
            "一级标签": path[0] if len(path) > 0 else "",
            "二级标签": path[1] if len(path) > 1 else "",
            "三级标签": path[2] if len(path) > 2 else "",
            "数量": count, "占比": count / max(total, 1),
        })
    return rows


def make_top_problems(clustered: list[dict[str, Any]], top_n: int = 20) -> list[dict[str, Any]]:
    groups = defaultdict(list)
    for row in clustered:
        groups[str(row.get("cluster_id", "unknown"))].append(row)
    rows = []
    total = len(clustered)
    for cid, items in sorted(groups.items(), key=lambda x: -len(x[1]))[:top_n]:
        points = Counter(p for x in items for p in (x.get("key_points") or []))
        reps = [x["cleaned_text"] for x in items[:3]]
        rows.append({"聚类ID": cid, "数量": len(items), "占比": len(items) / max(total, 1), "高频问题点": "；".join(k for k, _ in points.most_common(5)), "代表原声": " | ".join(reps)})
    return rows


def coverage_suggestions(distribution: list[dict[str, Any]], high: float, low: float, min_count: int) -> list[dict[str, Any]]:
    out = []
    for row in distribution:
        share, count = float(row["占比"]), int(row["数量"])
        path = "/".join(x for x in [row["一级标签"], row["二级标签"], row["三级标签"]] if x)
        if path == "待复核/未知":
            action = "重新聚类未知样本并检查标签缺口" if share > 0.10 else "持续观察"
        elif share > high:
            action = "检查簇内稳定子结构，满足纯度与业务价值后拆分"
        elif share < low and count < min_count:
            action = "检查是否高风险/新兴问题；否则合并或转观察项"
        else:
            action = "保持"
        out.append({"标签路径": path, "数量": count, "占比": share, "治理建议": action})
    return out


def label_and_report(args: argparse.Namespace) -> dict[str, Any]:
    work = Path(args.work_dir).expanduser().resolve(); work.mkdir(parents=True, exist_ok=True)
    records = read_jsonl(work / "normalized_records.jsonl")
    features = read_jsonl(Path(args.features).expanduser().resolve())
    taxonomy = load_json(Path(args.taxonomy).expanduser().resolve())
    validation = validate_taxonomy_obj(taxonomy, args.embedding_model, args.sibling_similarity_threshold, args.max_children)
    dump_json(work / "taxonomy_validation.json", validation)
    if not validation["valid"]:
        raise ValueError("标签体系未通过结构校验")
    by_feature = {str(x["record_id"]): x for x in features}
    errors = validate_features(records, features)
    if errors:
        dump_json(work / "feature_validation_errors.json", errors); raise ValueError("语义特征未通过校验")
    labels, children = build_taxonomy(taxonomy)
    if args.classifier == "laya":
        classifier = LayaClassifier(args.laya_model, args.laya_subfolder)
        model_version = f"laya:{Path(args.laya_model).name}/{args.laya_subfolder}"
    else:
        classifier = EmbeddingClassifier(args.embedding_model)
        model_version = f"embedding:{Path(args.embedding_model).name}"
    tagged = []
    for r in records:
        f = by_feature[str(r["record_id"])]
        for issue_idx, text, evidence in issue_texts(r, f):
            result = hierarchical_label(text, classifier, labels, children, args.confidence_threshold, args.margin_threshold)
            tagged.append({
                "record_id": r["record_id"], "source_row": r["source_row"], "issue_index": issue_idx,
                "raw_text": r["raw_text"], "cleaned_text": r["cleaned_text"], "evidence": evidence,
                "key_points": f.get("key_points", []), "problem_object": f.get("problem_object", ""),
                "user_request": f.get("user_request", ""), "emotion": f.get("emotion", ""),
                "severity": f.get("severity", ""), "taxonomy_version": taxonomy.get("version", "unknown"),
                "model_version": model_version, **result,
            })
    write_jsonl(work / "tagged_data.jsonl", tagged)
    review = [x for x in tagged if x["status"] != "labeled"]
    write_jsonl(work / "review_queue.jsonl", review)
    distribution = make_distribution(tagged)
    dump_json(work / "label_distribution.json", distribution)
    clustered_path = work / "clustered_records.jsonl"
    clustered = read_jsonl(clustered_path) if clustered_path.exists() else []
    top = make_top_problems(clustered)
    dump_json(work / "top_problems.json", top)
    suggestions = coverage_suggestions(distribution, args.coverage_high, args.coverage_low, args.min_label_count)
    dump_json(work / "coverage_suggestions.json", suggestions)
    training_report_path = Path(args.laya_model).expanduser().resolve() / "training_report.json" if args.classifier == "laya" and args.laya_model else None
    training_report = load_json(training_report_path) if training_report_path and training_report_path.exists() else None
    manifest = {
        "generated_at": time.strftime("%Y-%m-%d %H:%M:%S"), "record_count": len(records),
        "issue_count": len(tagged), "labeled_count": len(tagged) - len(review), "review_count": len(review),
        "review_rate": round(len(review) / max(len(tagged), 1), 4), "taxonomy_version": taxonomy.get("version"),
        "classifier": args.classifier, "model_version": model_version,
        "similarity_threshold": (load_json(work / "cluster_manifest.json").get("selected_similarity_threshold") if (work / "cluster_manifest.json").exists() else args.similarity_threshold),
        "confidence_threshold": args.confidence_threshold, "margin_threshold": args.margin_threshold,
        "embedding_model": args.embedding_model, "taxonomy_validation": validation,
        "training_report": training_report,
    }
    dump_json(work / "run_manifest.json", manifest)
    if args.output_xlsx:
        export_workbook(work, taxonomy, tagged, distribution, top, review, suggestions, Path(args.output_xlsx).expanduser().resolve())
    return manifest


def scalar(value: Any) -> Any:
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    if isinstance(value, np.generic):
        return value.item()
    return value


def write_sheet(ws, headers: list[str], rows: list[dict[str, Any]]) -> None:
    from copy import copy
    from openpyxl.styles import Alignment, Font, PatternFill, Side
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.utils import get_column_letter
    ws.append(headers)
    for row in rows:
        ws.append([scalar(row.get(h, "")) for h in headers])
    header_fill = PatternFill("solid", fgColor="4472C4")
    line = Side(style="thin", color="B8C7D9")
    for cell in ws[1]:
        cell.font = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        border = copy(cell.border); border.bottom = line; cell.border = border
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="微软雅黑", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions; ws.sheet_view.showGridLines = False
    for idx, h in enumerate(headers, 1):
        vals = [len(str(ws.cell(r, idx).value or "")) for r in range(1, min(ws.max_row, 100) + 1)]
        width = min(max(max(vals, default=8) + 2, 10), 45)
        ws.column_dimensions[get_column_letter(idx)].width = width
    if ws.max_row >= 2 and ws.max_column >= 1:
        digest = hashlib.sha1(ws.title.encode("utf-8")).hexdigest()[:10]
        display = f"T_{digest}"
        table = Table(displayName=display, ref=ws.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=False, showColumnStripes=False)
        ws.add_table(table)


def export_workbook(work: Path, taxonomy: dict[str, Any], tagged: list[dict[str, Any]], distribution: list[dict[str, Any]], top: list[dict[str, Any]], review: list[dict[str, Any]], suggestions: list[dict[str, Any]], output: Path) -> None:
    from openpyxl import Workbook
    wb = Workbook(); wb.remove(wb.active)
    manifest = load_json(work / "run_manifest.json") if (work / "run_manifest.json").exists() else {}
    profile = load_json(work / "data_profile.json") if (work / "data_profile.json").exists() else {}
    diagnostics = load_json(work / "threshold_diagnostics.json") if (work / "threshold_diagnostics.json").exists() else []
    instructions = [
        {"项目": "运行状态", "内容": "完成"},
        {"项目": "生成时间", "内容": manifest.get("generated_at", "")},
        {"项目": "原始记录数", "内容": profile.get("raw_count", "")},
        {"项目": "有效记录数", "内容": profile.get("valid_count", manifest.get("record_count", ""))},
        {"项目": "问题点数", "内容": manifest.get("issue_count", "")},
        {"项目": "人工复核率", "内容": manifest.get("review_rate", "")},
        {"项目": "标签版本", "内容": manifest.get("taxonomy_version", "")},
        {"项目": "分类器", "内容": manifest.get("model_version", "")},
        {"项目": "向量模型", "内容": manifest.get("embedding_model", "")},
        {"项目": "语义相似度阈值", "内容": manifest.get("similarity_threshold", "")},
        {"项目": "分类置信阈值", "内容": manifest.get("confidence_threshold", "")},
        {"项目": "Top1-Top2 最小差值", "内容": manifest.get("margin_threshold", "")},
        {"项目": "训练与校准", "内容": (f"已执行；验证集 Macro-F1={manifest['training_report'].get('after_calibrated', {}).get('macro_f1')}，ECE={manifest['training_report'].get('after_calibrated', {}).get('ece')}，温度={manifest['training_report'].get('temperature')}" if manifest.get("training_report") else "未提供人工金标，当前为基础模型候选结果，正式使用前需金标评估与校准")},
        {"项目": "处理说明", "内容": "执行 AI 完成特征抽取、候选标签生成与 MECE 评审；脚本完成数据清洗、BGE 向量聚类、Laya 分层打标、拒识、质量统计和结果落表。"},
    ]
    for d in diagnostics:
        instructions.append({"项目": f"阈值诊断 {d['similarity_threshold']}", "内容": f"簇数={d['cluster_count']}，噪声率={d['noise_rate']}，最大簇占比={d['largest_cluster_share']}，轮廓系数={d['silhouette_cosine']}"})
    ws = wb.create_sheet("运行说明"); write_sheet(ws, ["项目", "内容"], instructions)
    detail_rows = []
    for x in tagged:
        detail_rows.append({
            "record_id": x["record_id"], "source_row": x["source_row"], "issue_index": x["issue_index"],
            "用户原声": x["raw_text"], "证据片段": x["evidence"], "关键问题点": "；".join(x.get("key_points", [])),
            "问题对象": x.get("problem_object", ""), "用户诉求": x.get("user_request", ""), "情绪": x.get("emotion", ""),
            "严重度": x.get("severity", ""), "一级标签": x["label_path"][0] if len(x["label_path"]) > 0 else "",
            "二级标签": x["label_path"][1] if len(x["label_path"]) > 1 else "", "三级标签": x["label_path"][2] if len(x["label_path"]) > 2 else "",
            "标签ID路径": "/".join(x["label_ids"]), "层级概率": "/".join(map(str, x["probabilities"])),
            "综合置信度": x["confidence"], "状态": x["status"], "复核原因": x["review_reason"] or "",
            "标签版本": x["taxonomy_version"], "模型版本": x["model_version"],
        })
    ws = wb.create_sheet("打标明细"); write_sheet(ws, list(detail_rows[0]) if detail_rows else ["record_id"], detail_rows)
    tax_rows = []
    for x in taxonomy.get("labels", []):
        tax_rows.append({
            "label_id": x["label_id"], "parent_id": x.get("parent_id") or "", "层级": x["level"], "标签名": x["name"],
            "定义": x["definition"], "包含": "；".join(x.get("include", [])), "不包含": "；".join(x.get("exclude", [])),
            "正例": " | ".join(x.get("positive_examples", [])), "反例": " | ".join(x.get("negative_examples", [])),
            "业务动作": x.get("business_action", ""), "状态": x.get("status", "active"), "版本": taxonomy.get("version", ""),
        })
    ws = wb.create_sheet("标签体系"); write_sheet(ws, list(tax_rows[0]) if tax_rows else ["label_id"], tax_rows)
    dist_rows = [{**x, **next(({"治理建议": y["治理建议"]} for y in suggestions if y["标签路径"] == "/".join(v for v in [x["一级标签"], x["二级标签"], x["三级标签"]] if v)), {"治理建议": ""})} for x in distribution]
    ws = wb.create_sheet("标签分布"); write_sheet(ws, ["一级标签", "二级标签", "三级标签", "数量", "占比", "治理建议"], dist_rows)
    dist_end = ws.max_row
    for row in range(2, dist_end + 1):
        ws.cell(row=row, column=5, value=f"=IF(SUM($D$2:$D${dist_end})=0,0,D{row}/SUM($D$2:$D${dist_end}))")
        ws.cell(row=row, column=5).number_format = "0.00%"
    ws = wb.create_sheet("TOP问题"); write_sheet(ws, ["聚类ID", "数量", "占比", "高频问题点", "代表原声"], top)
    top_end = ws.max_row
    for row in range(2, top_end + 1):
        ws.cell(row=row, column=3, value=f"=IF(SUM($B$2:$B${top_end})=0,0,B{row}/SUM($B$2:$B${top_end}))")
        ws.cell(row=row, column=3).number_format = "0.00%"
    review_rows = [{"record_id": x["record_id"], "issue_index": x["issue_index"], "用户原声": x["raw_text"], "已判路径": "/".join(x["label_path"]), "综合置信度": x["confidence"], "复核原因": x["review_reason"], "人工标签": "", "复核备注": ""} for x in review]
    ws = wb.create_sheet("人工复核队列"); write_sheet(ws, ["record_id", "issue_index", "用户原声", "已判路径", "综合置信度", "复核原因", "人工标签", "复核备注"], review_rows)
    output.parent.mkdir(parents=True, exist_ok=True); wb.save(output)


def distribution_map(rows: list[dict[str, Any]]) -> dict[str, float]:
    out = {}
    for row in rows:
        path = "/".join(x for x in [row.get("一级标签", ""), row.get("二级标签", ""), row.get("三级标签", "")] if x)
        out[path] = float(row.get("占比", 0.0))
    return out


def review_runs(args: argparse.Namespace) -> dict[str, Any]:
    current_dir = Path(args.current_dir).expanduser().resolve()
    previous_dir = Path(args.previous_dir).expanduser().resolve() if args.previous_dir else None
    current_dist = load_json(current_dir / "label_distribution.json")
    current_manifest = load_json(current_dir / "run_manifest.json")
    report: dict[str, Any] = {
        "current_taxonomy_version": current_manifest.get("taxonomy_version"),
        "current_review_rate": current_manifest.get("review_rate"),
        "unknown_alert": float(current_manifest.get("review_rate", 0)) > args.unknown_alert,
        "label_drift": [], "js_divergence": None, "new_labels": [], "missing_labels": [],
    }
    if previous_dir:
        prev_dist = load_json(previous_dir / "label_distribution.json")
        prev_manifest = load_json(previous_dir / "run_manifest.json")
        cur, prev = distribution_map(current_dist), distribution_map(prev_dist)
        keys = sorted(set(cur) | set(prev))
        p = np.asarray([prev.get(k, 0.0) for k in keys], dtype=float)
        q = np.asarray([cur.get(k, 0.0) for k in keys], dtype=float)
        p = p / max(p.sum(), 1e-12); q = q / max(q.sum(), 1e-12); m = (p + q) / 2
        def kl(a, b):
            mask = a > 0
            return float(np.sum(a[mask] * np.log2(a[mask] / np.clip(b[mask], 1e-12, None))))
        report["js_divergence"] = round((kl(p, m) + kl(q, m)) / 2, 6)
        report["previous_taxonomy_version"] = prev_manifest.get("taxonomy_version")
        report["review_rate_delta"] = round(float(current_manifest.get("review_rate", 0)) - float(prev_manifest.get("review_rate", 0)), 4)
        report["new_labels"] = sorted(set(cur) - set(prev)); report["missing_labels"] = sorted(set(prev) - set(cur))
        for k in keys:
            delta = q[keys.index(k)] - p[keys.index(k)]
            if abs(delta) >= args.share_delta_alert:
                report["label_drift"].append({"label_path": k, "previous_share": round(float(p[keys.index(k)]), 4), "current_share": round(float(q[keys.index(k)]), 4), "delta": round(float(delta), 4)})
    dump_json(Path(args.output or current_dir / "review_report.json"), report)
    return report


def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(description="用户原声标签发现、治理、Laya 打标与报表流水线")
    sub = p.add_subparsers(dest="command", required=True)
    s = sub.add_parser("prepare")
    s.add_argument("--input", required=True); s.add_argument("--text-col", required=True); s.add_argument("--id-col")
    s.add_argument("--context-cols", nargs="*", default=[]); s.add_argument("--sheet-name"); s.add_argument("--output-dir", required=True)
    s.add_argument("--min-chars", type=int, default=4); s.add_argument("--batch-size", type=int, default=50)
    s.set_defaults(func=prepare_records)
    s = sub.add_parser("cluster")
    s.add_argument("--work-dir", required=True); s.add_argument("--features", required=True); s.add_argument("--embedding-model", required=True)
    s.add_argument("--similarity-threshold", type=float, default=None, help="省略时从 threshold-grid 自动选择")
    s.add_argument("--threshold-grid", nargs="*", type=float, default=DEFAULT_THRESHOLDS)
    s.add_argument("--min-cluster-size", type=int, default=2); s.add_argument("--embedding-batch-size", type=int, default=32)
    s.set_defaults(func=cluster_records)
    s = sub.add_parser("validate-taxonomy")
    s.add_argument("--taxonomy", required=True); s.add_argument("--embedding-model"); s.add_argument("--output")
    s.add_argument("--sibling-similarity-threshold", type=float, default=0.86); s.add_argument("--max-children", type=int, default=20)
    s.set_defaults(func=validate_taxonomy_cmd)
    s = sub.add_parser("label-report")
    s.add_argument("--work-dir", required=True); s.add_argument("--features", required=True); s.add_argument("--taxonomy", required=True); s.add_argument("--output-xlsx", help="可选故障兜底；标准交付由执行 AI 调用 sheet Skill 直接写飞书表格")
    s.add_argument("--classifier", choices=["laya", "embedding"], default="laya"); s.add_argument("--laya-model"); s.add_argument("--laya-subfolder", default="multilingual")
    s.add_argument("--embedding-model", required=True); s.add_argument("--similarity-threshold", type=float, default=0.72)
    s.add_argument("--confidence-threshold", type=float, default=0.50); s.add_argument("--margin-threshold", type=float, default=0.10)
    s.add_argument("--sibling-similarity-threshold", type=float, default=0.86); s.add_argument("--max-children", type=int, default=20)
    s.add_argument("--coverage-high", type=float, default=0.25); s.add_argument("--coverage-low", type=float, default=0.005); s.add_argument("--min-label-count", type=int, default=20)
    s.set_defaults(func=label_and_report)
    s = sub.add_parser("review")
    s.add_argument("--current-dir", required=True); s.add_argument("--previous-dir"); s.add_argument("--output")
    s.add_argument("--unknown-alert", type=float, default=0.10); s.add_argument("--share-delta-alert", type=float, default=0.05)
    s.set_defaults(func=review_runs)
    return p


def main() -> None:
    args = build_parser().parse_args()
    if getattr(args, "classifier", None) == "laya" and not args.laya_model:
        raise SystemExit("classifier=laya 时必须提供 --laya-model")
    result = args.func(args)
    print(json.dumps({"ok": True, "command": args.command, "result": result}, ensure_ascii=False))


if __name__ == "__main__":
    main()
